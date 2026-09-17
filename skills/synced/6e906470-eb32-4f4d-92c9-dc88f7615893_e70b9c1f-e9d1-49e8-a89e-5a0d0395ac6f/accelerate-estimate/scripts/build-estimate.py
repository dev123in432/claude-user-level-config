#!/usr/bin/env python3
"""Accelerate Tech — Project Estimation Workbook Generator
Supports: sprint, phase, t&m, managed_services models

Formula-driven outputs: inputs (Days, rate-card values, Contingency %, GST %,
Discount %, milestone %) are values; everything downstream (Daily Rate via
VLOOKUP, Estimate, Estimate Inc. Contingency, totals, Discount, Subtotal,
GST, Inc GST, Resource Summary rows, Weekly View row totals + grand totals,
Forecast Export Estimated Cost, Payment Milestones Amount Ex/Inc GST,
Hourly Rate) are formulas. Edit any input cell and totals recalculate on
workbook open.

Known limitation: Weekly View per-cell allocations remain values because they
are derived from working-day math against phase/sprint date ranges. Row Total
Days, Total Cost and the grand totals at the bottom ARE formulas, so they
recalculate if you tweak an individual cell."""

import json, sys, math
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from calendar import monthrange
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Branding ──────────────────────────────────────────────────────────────────
NAVY = "19263C"
TEAL = "6BE1B8"
WHITE = "FFFFFF"
LIGHT_GREY = "F2F2F2"
AMBER = "FFF3CD"
RED_BG = "F8D7DA"
GREEN_BG = "D4EDDA"

navy_fill = PatternFill("solid", fgColor=NAVY)
teal_fill = PatternFill("solid", fgColor=TEAL)
grey_fill = PatternFill("solid", fgColor=LIGHT_GREY)
green_fill = PatternFill("solid", fgColor=GREEN_BG)
amber_fill = PatternFill("solid", fgColor=AMBER)
red_fill = PatternFill("solid", fgColor=RED_BG)

hdr_font = Font(name="Arial", size=10, bold=True, color=WHITE)
teal_font = Font(name="Arial", size=10, bold=True, color="000000")
body_font = Font(name="Arial", size=10)
bold_font = Font(name="Arial", size=10, bold=True)
title_font = Font(name="Arial", size=14, bold=True, color=NAVY)
subtitle_font = Font(name="Arial", size=11, bold=True, color=NAVY)
note_font = Font(name="Arial", size=9, italic=True, color="666666")

thin_border = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)

CURRENCY_FMT = '_-"$"* #,##0_-;\\-"$"* #,##0_-;_-"$"* "-"??_-;_-@_-'
PCT_FMT = "0%"
DEC_FMT = "0.0"

# ── Rate Cards & Defaults ─────────────────────────────────────────────────────
RATE_CARDS = {
    "FY27": {"Principal": 2228, "Senior": 2022, "Associate": 1730, "Junior": 1494},
    "FY26": {"Principal": 2163, "Senior": 1963, "Associate": 1680, "Junior": 1450},
    "FY25": {"Principal": 2035, "Senior": 1870, "Associate": 1595, "Junior": 1320},
}

CONTINGENCY_BY_AREA = {
    "Project Management": 0.10,
    "Setup and Design": 0.10,
    "Development": 0.15,
    "Integration": 0.20,
    "Data Migration": 0.20,
    "Testing": 0.15,
    "UAT Support": 0.15,
    "Go-Live / Cutover": 0.15,
    "Handover / Transition": 0.10,
    "Training": 0.10,
    "Managed Services": 0.10,
}

MODEL_LABELS = {
    "sprint": "Fixed Price (Sprint-Based)",
    "phase": "Fixed Price (Days-Per-Phase)",
    "t&m": "Time & Materials",
    "managed_services": "Managed Services",
}


# ── Layout dataclass — tracks cell positions so formulas can reference them ──
@dataclass
class Layout:
    # Config sheet anchors
    gst_cell: str = ""               # e.g. "Config!$B$10"
    discount_cell: str = ""          # e.g. "Config!$B$11" or "" if none
    rate_card_range: str = ""        # e.g. "Config!$A$18:$B$21"
    rate_card_band_col: str = "A"    # column letter for Rate Band in rate card
    rate_card_rate_col: str = "B"    # column letter for Daily Rate in rate card
    contingency_range: str = ""      # e.g. "Config!$A$25:$B$35" or ""
    # Estimate sheet anchors
    est_first_row: int = 2
    est_last_row: int = 2
    est_days_col: int = 0
    est_rate_col: int = 0
    est_estimate_col: int = 0
    est_cont_pct_col: int = 0
    est_est_cont_col: int = 0
    est_func_col: int = 0
    est_area_col: int = 0
    est_rateband_col: int = 0
    est_total_row: int = 0           # row number of TOTAL row
    est_discount_row: int = 0        # row number of Discount row (0 if none)
    est_subtotal_row: int = 0        # row number of Subtotal After Discount row
    est_gst_row: int = 0             # row number of GST row
    est_inc_gst_row: int = 0         # row number of Total Inc. GST row
    est_n_cols: int = 0
    # The "fee basis" cell that milestones reference
    # For fixed price = Total Inc. Contingency on TOTAL row; T&M = Estimate on TOTAL row
    est_fee_basis_ex_gst_cell: str = ""
    est_fee_basis_inc_gst_cell: str = ""


# ── Helpers ───────────────────────────────────────────────────────────────────
def col_letter(n):
    return get_column_letter(n)


def sheet_ref(sheet_name, cell):
    """Return a sheet-qualified cell reference; single-quote sheet names with spaces."""
    if " " in sheet_name:
        return f"'{sheet_name}'!{cell}"
    return f"{sheet_name}!{cell}"


def style_header_row(ws, row, max_col, fill=None, font=None):
    for c in range(1, max_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill or navy_fill
        cell.font = font or hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border


def style_total_row(ws, row, max_col, fill=None, font=None):
    for c in range(1, max_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill or teal_fill
        cell.font = font or teal_font
        cell.border = thin_border


def write_row(ws, row, values, start_col=1, fmt_map=None):
    """Write VALUE cells (inputs / display-only)."""
    for i, v in enumerate(values):
        col = start_col + i
        cell = ws.cell(row=row, column=col, value=v)
        cell.font = body_font
        cell.border = thin_border
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if fmt_map and col in fmt_map:
            cell.number_format = fmt_map[col]


def write_cell(ws, row, col, value, fmt=None, font=None, fill=None):
    """Generic single-cell writer that styles consistently. `value` can be a formula
    string starting with '=' — openpyxl treats that as a formula automatically."""
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font or body_font
    cell.border = thin_border
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    return cell


def set_col_widths(ws, widths):
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = w


def get_discount(config):
    model = config.get("model", "sprint")
    if model in ("t&m", "managed_services"):
        return 0
    return config.get("discount", 0)


def get_contingency(area, config):
    model = config.get("model", "sprint")
    if model in ("t&m", "managed_services"):
        return 0
    flat = config.get("flatContingency")
    if flat is not None:
        return flat
    return CONTINGENCY_BY_AREA.get(area, 0.15)


def get_task_days(task, config):
    if config["model"] == "sprint":
        return task.get("allocation", 0) * config.get("sprintLength", 10)
    elif config["model"] in ("phase", "t&m"):
        return task.get("days", 0)
    return 0


def add_month(d, n):
    month = d.month - 1 + n
    year = d.year + month // 12
    month = month % 12 + 1
    day = min(d.day, monthrange(year, month)[1])
    return d.replace(year=year, month=month, day=day)


# ── Weekly data for all models ────────────────────────────────────────────────
def get_project_weeks(config):
    model = config["model"]
    if model == "sprint":
        all_dates = []
        for s in config["sprints"]:
            all_dates.append(datetime.strptime(s["startDate"], "%Y-%m-%d"))
            all_dates.append(datetime.strptime(s["endDate"], "%Y-%m-%d"))
    elif model in ("phase", "t&m"):
        all_dates = []
        for p in config["phases"]:
            all_dates.append(datetime.strptime(p["startDate"], "%Y-%m-%d"))
            all_dates.append(datetime.strptime(p["endDate"], "%Y-%m-%d"))
    elif model == "managed_services":
        start = datetime.strptime(config["startDate"], "%Y-%m-%d")
        months = config.get("contractMonths", 12)
        end = add_month(start, months) - timedelta(days=1)
        all_dates = [start, end]
    else:
        return []

    min_d = min(all_dates)
    max_d = max(all_dates)
    start_mon = min_d - timedelta(days=min_d.weekday())
    end_mon = max_d - timedelta(days=max_d.weekday())
    weeks = []
    cur = start_mon
    while cur <= end_mon:
        weeks.append(cur)
        cur += timedelta(days=7)
    return weeks


def build_weekly_data(config, rates):
    weeks = get_project_weeks(config)
    model = config["model"]

    if model == "managed_services":
        allocs = config.get("monthlyAllocations", [])
        functions = sorted(set(a["function"] for a in allocs))
    else:
        functions = sorted(set(t["function"] for t in config.get("tasks", [])))

    weekly = {w: {f: 0.0 for f in functions} for w in weeks}
    weekly_cost = {w: {f: 0.0 for f in functions} for w in weeks}
    # Track which rate band each function uses for the rate row formula
    func_rateband = {}

    if model == "sprint":
        sprint_map = {s["id"]: s for s in config["sprints"]}
        for task in config["tasks"]:
            func_rateband.setdefault(task["function"], task.get("rateBand", "Senior"))
            sp = sprint_map[task["sprint"]]
            sp_start = datetime.strptime(sp["startDate"], "%Y-%m-%d")
            sp_end = datetime.strptime(sp["endDate"], "%Y-%m-%d")
            alloc = task["allocation"]
            sprint_len = config.get("sprintLength", 10)
            total_days = alloc * sprint_len
            working_days = sum(1 for d in (sp_start + timedelta(n) for n in range((sp_end - sp_start).days + 1)) if d.weekday() < 5)
            if working_days == 0:
                continue
            daily_alloc = total_days / working_days
            rate = rates.get(task.get("rateBand", "Senior"), 1963)
            cur = sp_start
            while cur <= sp_end:
                if cur.weekday() < 5:
                    wk = cur - timedelta(days=cur.weekday())
                    if wk in weekly:
                        weekly[wk][task["function"]] += daily_alloc
                        weekly_cost[wk][task["function"]] += daily_alloc * rate
                cur += timedelta(days=1)

    elif model in ("phase", "t&m"):
        for task in config["tasks"]:
            func_rateband.setdefault(task["function"], task.get("rateBand", "Senior"))
            phase = next(p for p in config["phases"] if p["name"] == task["phase"])
            p_start = datetime.strptime(phase["startDate"], "%Y-%m-%d")
            p_end = datetime.strptime(phase["endDate"], "%Y-%m-%d")
            working_days = sum(1 for d in (p_start + timedelta(n) for n in range((p_end - p_start).days + 1)) if d.weekday() < 5)
            if working_days == 0:
                continue
            daily_alloc = task["days"] / working_days
            rate = rates.get(task.get("rateBand", "Senior"), 1963)
            cur = p_start
            while cur <= p_end:
                if cur.weekday() < 5:
                    wk = cur - timedelta(days=cur.weekday())
                    if wk in weekly:
                        weekly[wk][task["function"]] += daily_alloc
                        weekly_cost[wk][task["function"]] += daily_alloc * rate
                cur += timedelta(days=1)

    elif model == "managed_services":
        start = datetime.strptime(config["startDate"], "%Y-%m-%d")
        months = config.get("contractMonths", 12)
        annual_uplift = config.get("annualUplift", 0)
        for alloc_item in allocs:
            fn = alloc_item["function"]
            func_rateband.setdefault(fn, alloc_item.get("rateBand", "Senior"))
            rb = alloc_item.get("rateBand", "Senior")
            base_rate = rates.get(rb, 1963)
            dpm = alloc_item["daysPerMonth"]
            for m in range(months):
                m_start = add_month(start, m)
                m_end = add_month(start, m + 1) - timedelta(days=1)
                year_idx = m // 12
                rate = base_rate * ((1 + annual_uplift) ** year_idx)
                wd_in_month = sum(1 for d in (m_start + timedelta(n) for n in range((m_end - m_start).days + 1)) if d.weekday() < 5)
                if wd_in_month == 0:
                    continue
                daily_alloc = dpm / wd_in_month
                cur = m_start
                while cur <= m_end:
                    if cur.weekday() < 5:
                        wk = cur - timedelta(days=cur.weekday())
                        if wk in weekly:
                            weekly[wk][fn] += daily_alloc
                            weekly_cost[wk][fn] += daily_alloc * rate
                    cur += timedelta(days=1)

    return weeks, functions, weekly, weekly_cost, func_rateband


def get_phase_for_week(config, week_date):
    model = config["model"]
    if model == "sprint":
        phase_map = {}
        for p in config.get("phases", []):
            for sid in p.get("sprints", []):
                phase_map[sid] = p["name"]
        for s in config["sprints"]:
            s_start = datetime.strptime(s["startDate"], "%Y-%m-%d")
            s_end = datetime.strptime(s["endDate"], "%Y-%m-%d")
            if s_start <= week_date + timedelta(days=4) and s_end >= week_date:
                return phase_map.get(s["id"], "")
    elif model in ("phase", "t&m"):
        for p in config["phases"]:
            p_start = datetime.strptime(p["startDate"], "%Y-%m-%d")
            p_end = datetime.strptime(p["endDate"], "%Y-%m-%d")
            if p_start <= week_date + timedelta(days=4) and p_end >= week_date:
                return p["name"]
    elif model == "managed_services":
        mid_week = week_date + timedelta(days=2)
        return mid_week.strftime("%B %Y")
    return ""


# ── Config Sheet ──────────────────────────────────────────────────────────────
def build_config_sheet(wb, config, rates, layout: Layout):
    ws = wb.active
    ws.title = "Config"
    ws.sheet_properties.tabColor = NAVY
    set_col_widths(ws, [28, 22, 18, 22, 18])

    model = config.get("model", "sprint")
    ws["A1"] = "Project Estimation"
    ws["A1"].font = title_font

    # Inputs banner (row 2)
    ws["A2"] = ("Inputs are editable; downstream cells are formulas. "
                "Change a Daily Rate, Days, Contingency %, GST % or Discount % and totals recalculate.")
    ws["A2"].font = note_font
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=5)

    fields = [
        ("Project Name", config.get("projectName", "")),
        ("Client", config.get("clientName", "")),
        ("Version", config.get("estimateVersion", "V1")),
        ("Date", config.get("estimateDate", "")),
        ("Prepared By", config.get("preparedBy", "")),
        ("Model", MODEL_LABELS.get(model, model)),
        ("Financial Year", config.get("financialYear", "FY26")),
        ("GST", config.get("gst", 0.10)),
    ]
    gst_field_idx = len(fields) - 1  # GST is the last in this list

    discount_field_idx = None
    if model in ("sprint", "phase"):
        discount_field_idx = len(fields)
        fields.append(("Discount", config.get("discount", 0)))
    if model == "t&m":
        cap = config.get("budgetCap", 0)
        cap_label = "Budget Cap (Inc GST)" if config.get("budgetCapInclGST", False) else "Budget Cap (Ex GST)"
        fields.append((cap_label, cap))
        fields.append(("Invoice Frequency", config.get("invoiceFrequency", "Monthly").title()))
    if model == "managed_services":
        fields.append(("Contract Months", config.get("contractMonths", 12)))
        fields.append(("Start Date", config.get("startDate", "")))
        fields.append(("Annual Uplift", config.get("annualUplift", 0)))

    r = 3  # start after banner
    for idx, (lbl, val) in enumerate(fields):
        ws.cell(row=r, column=1, value=lbl).font = bold_font
        c = ws.cell(row=r, column=2, value=val)
        c.font = body_font
        if isinstance(val, float) and val < 1:
            c.number_format = PCT_FMT
        elif isinstance(val, (int, float)) and val > 100:
            c.number_format = CURRENCY_FMT
        # Capture GST and Discount cell references
        if idx == gst_field_idx:
            layout.gst_cell = f"Config!$B${r}"
        if discount_field_idx is not None and idx == discount_field_idx:
            layout.discount_cell = f"Config!$B${r}"
        r += 1

    # Rate card
    r += 1
    ws.cell(row=r, column=1, value="Rate Card").font = subtitle_font
    r += 1
    for i, h in enumerate(["Rate Band", "Daily Rate", "Hourly Rate"]):
        ws.cell(row=r, column=1 + i, value=h)
    style_header_row(ws, r, 3)
    r += 1
    rate_card_start = r
    for band in ["Principal", "Senior", "Associate", "Junior"]:
        daily = rates[band]
        ws.cell(row=r, column=1, value=band).font = body_font
        ws.cell(row=r, column=1).border = thin_border
        # Daily Rate is the INPUT — keep as value
        c = ws.cell(row=r, column=2, value=daily)
        c.number_format = CURRENCY_FMT
        c.font = body_font
        c.border = thin_border
        # Hourly Rate is DERIVED from Daily Rate
        c = ws.cell(row=r, column=3, value=f"=B{r}/8")
        c.number_format = CURRENCY_FMT
        c.font = body_font
        c.border = thin_border
        r += 1
    rate_card_end = r - 1
    layout.rate_card_range = f"Config!$A${rate_card_start}:$B${rate_card_end}"

    # Contingency table (fixed price only)
    if model in ("sprint", "phase"):
        r += 1
        ws.cell(row=r, column=1, value="Risk-Weighted Contingency").font = subtitle_font
        r += 1
        ws.cell(row=r, column=1, value="Solution Area")
        ws.cell(row=r, column=2, value="Contingency %")
        style_header_row(ws, r, 2)
        r += 1
        flat = config.get("flatContingency")
        cont_start = r
        for area, pct in CONTINGENCY_BY_AREA.items():
            ws.cell(row=r, column=1, value=area).font = body_font
            ws.cell(row=r, column=1).border = thin_border
            c = ws.cell(row=r, column=2, value=flat if flat is not None else pct)
            c.number_format = PCT_FMT
            c.font = body_font
            c.border = thin_border
            r += 1
        cont_end = r - 1
        layout.contingency_range = f"Config!$A${cont_start}:$B${cont_end}"
        if flat is not None:
            ws.cell(row=r + 1, column=1, value=f"* Flat contingency override: {int(flat*100)}%").font = note_font
    elif model == "t&m":
        r += 1
        ws.cell(row=r, column=1, value="T&M Note").font = subtitle_font
        r += 1
        ws.cell(row=r, column=1, value="Contingency is 0% — risk sits with the client. Invoiced on actuals against rate card.").font = note_font
    elif model == "managed_services":
        r += 1
        ws.cell(row=r, column=1, value="Managed Services Note").font = subtitle_font
        r += 1
        ws.cell(row=r, column=1, value="Monthly retainer. Unused days do not roll over unless agreed. Annual uplift applied at each contract anniversary.").font = note_font


# ── Estimate Sheet (sprint/phase/t&m) ────────────────────────────────────────
def build_estimate_sheet(wb, config, rates, layout: Layout):
    model = config["model"]
    if model == "managed_services":
        return build_monthly_allocation_sheet(wb, config, rates, layout)

    ws = wb.create_sheet("Estimate")
    ws.sheet_properties.tabColor = NAVY
    is_sprint = model == "sprint"
    is_tm = model == "t&m"
    show_cont = model in ("sprint", "phase")

    # Inserted "Daily Rate" between "Rate Band" and "Solution Area"
    if is_sprint:
        headers = ["Phase", "Sprint", "Objective", "Function", "Rate Band", "Daily Rate",
                   "Solution Area", "Allocation", "Days", "Estimate"]
        col_widths = [18, 28, 55, 24, 14, 14, 22, 12, 10, 16]
        # column indices (1-based)
        c_phase, c_sprint, c_obj, c_func, c_rb, c_dr, c_sa, c_alloc, c_days, c_est = 1, 2, 3, 4, 5, 6, 7, 8, 9, 10
    else:
        headers = ["Phase", "Objective", "Function", "Rate Band", "Daily Rate",
                   "Solution Area", "Days", "Estimate"]
        col_widths = [18, 55, 24, 14, 14, 22, 10, 16]
        c_phase, c_obj, c_func, c_rb, c_dr, c_sa, c_days, c_est = 1, 2, 3, 4, 5, 6, 7, 8

    if show_cont:
        headers += ["Contingency %", "Estimate Inc. Contingency"]
        col_widths += [14, 22]
        c_cont = len(headers) - 1
        c_est_cont = len(headers)
    else:
        c_cont = 0
        c_est_cont = 0

    n_cols = len(headers)
    layout.est_n_cols = n_cols

    set_col_widths(ws, col_widths)
    for i, h in enumerate(headers):
        ws.cell(row=1, column=i + 1, value=h)
    style_header_row(ws, 1, n_cols)

    # Populate task rows
    r = 2
    first_task_row = r
    prev_phase = None
    sprint_map = {s["id"]: s for s in config.get("sprints", [])} if is_sprint else {}
    flat_cont = config.get("flatContingency") if show_cont else None

    for task in config["tasks"]:
        if is_sprint:
            sp = sprint_map.get(task["sprint"], {})
            phase = ""
            for p in config.get("phases", []):
                if task["sprint"] in p.get("sprints", []):
                    phase = p["name"]
                    break
            alloc = task.get("allocation", 0)
            days = alloc * config.get("sprintLength", 10)
        else:
            phase = task.get("phase", "")
            days = task.get("days", 0)

        show_phase = phase if phase != prev_phase else ""
        prev_phase = phase

        # Static value cells (text + Days input)
        if is_sprint:
            write_cell(ws, r, c_phase, show_phase, font=(bold_font if show_phase else body_font))
            write_cell(ws, r, c_sprint, sp.get("name", task["sprint"]))
            write_cell(ws, r, c_obj, task.get("objective", ""))
            write_cell(ws, r, c_func, task["function"])
            write_cell(ws, r, c_rb, task.get("rateBand", ""))
            # Daily Rate — formula lookup
            write_cell(ws, r, c_dr,
                       f"=VLOOKUP({col_letter(c_rb)}{r},{layout.rate_card_range},2,FALSE)",
                       fmt=CURRENCY_FMT)
            write_cell(ws, r, c_sa, task.get("solutionArea", ""))
            write_cell(ws, r, c_alloc, alloc, fmt=DEC_FMT)
            write_cell(ws, r, c_days, days, fmt=DEC_FMT)
        else:
            write_cell(ws, r, c_phase, show_phase, font=(bold_font if show_phase else body_font))
            write_cell(ws, r, c_obj, task.get("objective", ""))
            write_cell(ws, r, c_func, task["function"])
            write_cell(ws, r, c_rb, task.get("rateBand", ""))
            write_cell(ws, r, c_dr,
                       f"=VLOOKUP({col_letter(c_rb)}{r},{layout.rate_card_range},2,FALSE)",
                       fmt=CURRENCY_FMT)
            write_cell(ws, r, c_sa, task.get("solutionArea", ""))
            write_cell(ws, r, c_days, days, fmt=DEC_FMT)

        # Estimate = days * dailyRate
        write_cell(ws, r, c_est,
                   f"={col_letter(c_days)}{r}*{col_letter(c_dr)}{r}",
                   fmt=CURRENCY_FMT)

        if show_cont:
            # Contingency %: flat → value; per-area → VLOOKUP with IFERROR fallback
            if flat_cont is not None:
                write_cell(ws, r, c_cont, flat_cont, fmt=PCT_FMT)
            else:
                write_cell(ws, r, c_cont,
                           f"=IFERROR(VLOOKUP({col_letter(c_sa)}{r},{layout.contingency_range},2,FALSE),0.15)",
                           fmt=PCT_FMT)
            # Estimate Inc Contingency = Estimate * (1 + Contingency)
            write_cell(ws, r, c_est_cont,
                       f"={col_letter(c_est)}{r}*(1+{col_letter(c_cont)}{r})",
                       fmt=CURRENCY_FMT)

        r += 1

    last_task_row = r - 1

    # Capture layout
    layout.est_first_row = first_task_row
    layout.est_last_row = last_task_row
    layout.est_days_col = c_days
    layout.est_rate_col = c_dr
    layout.est_estimate_col = c_est
    layout.est_cont_pct_col = c_cont
    layout.est_est_cont_col = c_est_cont
    layout.est_func_col = c_func
    layout.est_area_col = c_sa
    layout.est_rateband_col = c_rb

    # Spacer
    r += 1

    # TOTAL row
    total_row = r
    layout.est_total_row = total_row
    ws.cell(row=r, column=1, value="TOTAL")
    style_total_row(ws, r, n_cols)

    days_L = col_letter(c_days)
    est_L = col_letter(c_est)
    write_cell(ws, r, c_days,
               f"=SUM({days_L}{first_task_row}:{days_L}{last_task_row})",
               fmt=DEC_FMT, font=teal_font, fill=teal_fill)
    write_cell(ws, r, c_est,
               f"=SUM({est_L}{first_task_row}:{est_L}{last_task_row})",
               fmt=CURRENCY_FMT, font=teal_font, fill=teal_fill)
    if show_cont:
        ec_L = col_letter(c_est_cont)
        write_cell(ws, r, c_est_cont,
                   f"=SUM({ec_L}{first_task_row}:{ec_L}{last_task_row})",
                   fmt=CURRENCY_FMT, font=teal_font, fill=teal_fill)

    # The "display" total for discount/GST cascade:
    # For sprint/phase with contingency → Inc Contingency column drives the cascade
    # For T&M → Estimate column drives the cascade
    # We need cells that reference TOTAL row figures.
    discount = get_discount(config)

    # Track which column drives discount/GST cascade
    cascade_col = c_est_cont if show_cont else c_est
    cascade_L = col_letter(cascade_col)
    total_ex_basis_cell = f"Estimate!${cascade_L}${total_row}"

    # Also keep the "Estimate" column for discount calc visualisation if present
    est_total_cell = f"Estimate!${est_L}${total_row}"
    est_cont_total_cell = f"Estimate!${col_letter(c_est_cont)}${total_row}" if show_cont else ""

    # Discount row
    if discount > 0 and layout.discount_cell:
        r += 1
        layout.est_discount_row = r
        ws.cell(row=r, column=cascade_col - 1, value="Discount").font = bold_font
        # Discount on the Estimate column
        write_cell(ws, r, c_est,
                   f"=-{est_L}{total_row}*{layout.discount_cell}",
                   fmt=CURRENCY_FMT)
        if show_cont:
            ec_L = col_letter(c_est_cont)
            write_cell(ws, r, c_est_cont,
                       f"=-{ec_L}{total_row}*{layout.discount_cell}",
                       fmt=CURRENCY_FMT)

        # Subtotal After Discount row
        r += 1
        layout.est_subtotal_row = r
        ws.cell(row=r, column=cascade_col - 1, value="Subtotal After Discount").font = bold_font
        write_cell(ws, r, c_est,
                   f"={est_L}{total_row}+{est_L}{layout.est_discount_row}",
                   fmt=CURRENCY_FMT)
        if show_cont:
            ec_L = col_letter(c_est_cont)
            write_cell(ws, r, c_est_cont,
                       f"={ec_L}{total_row}+{ec_L}{layout.est_discount_row}",
                       fmt=CURRENCY_FMT)
        # The "after discount" cell becomes the basis for GST
        basis_est_cell = f"{est_L}{layout.est_subtotal_row}"
        basis_cont_cell = f"{col_letter(c_est_cont)}{layout.est_subtotal_row}" if show_cont else ""
        basis_est_abs = f"${est_L}${layout.est_subtotal_row}"
        basis_cont_abs = f"${col_letter(c_est_cont)}${layout.est_subtotal_row}" if show_cont else ""
    else:
        basis_est_cell = f"{est_L}{total_row}"
        basis_cont_cell = f"{col_letter(c_est_cont)}{total_row}" if show_cont else ""
        basis_est_abs = f"${est_L}${total_row}"
        basis_cont_abs = f"${col_letter(c_est_cont)}${total_row}" if show_cont else ""

    # GST row
    r += 1
    layout.est_gst_row = r
    ws.cell(row=r, column=cascade_col - 1, value="GST").font = bold_font
    write_cell(ws, r, c_est,
               f"={basis_est_cell}*{layout.gst_cell}",
               fmt=CURRENCY_FMT)
    if show_cont:
        write_cell(ws, r, c_est_cont,
                   f"={basis_cont_cell}*{layout.gst_cell}",
                   fmt=CURRENCY_FMT)

    # Total Inc GST row
    r += 1
    layout.est_inc_gst_row = r
    style_total_row(ws, r, n_cols, fill=navy_fill, font=hdr_font)
    ws.cell(row=r, column=cascade_col - 1, value="Total Inc. GST").font = hdr_font
    write_cell(ws, r, c_est,
               f"={basis_est_cell}*(1+{layout.gst_cell})",
               fmt=CURRENCY_FMT, font=hdr_font, fill=navy_fill)
    if show_cont:
        write_cell(ws, r, c_est_cont,
                   f"={basis_cont_cell}*(1+{layout.gst_cell})",
                   fmt=CURRENCY_FMT, font=hdr_font, fill=navy_fill)

    # Set fee basis cells for milestones (absolute refs to survive any column shifts)
    # Fixed price (show_cont): basis = Inc-Contingency cell on subtotal-or-total row
    # T&M: basis = Estimate column on subtotal-or-total row
    if show_cont:
        layout.est_fee_basis_ex_gst_cell = f"Estimate!{basis_cont_abs}"
        layout.est_fee_basis_inc_gst_cell = f"Estimate!${col_letter(c_est_cont)}${layout.est_inc_gst_row}"
    else:
        layout.est_fee_basis_ex_gst_cell = f"Estimate!{basis_est_abs}"
        layout.est_fee_basis_inc_gst_cell = f"Estimate!${est_L}${layout.est_inc_gst_row}"

    # T&M budget cap comparison
    if is_tm and config.get("budgetCap"):
        r += 2
        cap = config["budgetCap"]
        cap_incl_gst = config.get("budgetCapInclGST", False)
        gst = config.get("gst", 0.10)
        cap_ex = cap / (1 + gst) if cap_incl_gst else cap
        # These remain values (budget cap is itself a config value)
        ws.cell(row=r, column=cascade_col - 1, value="Budget Cap (Ex GST)").font = bold_font
        c = ws.cell(row=r, column=c_est, value=cap_ex)
        c.number_format = CURRENCY_FMT
        r += 1
        ws.cell(row=r, column=cascade_col - 1, value="Remaining Budget").font = bold_font
        # Formula = cap_ex - basis
        c = ws.cell(row=r, column=c_est, value=f"={cap_ex}-{basis_est_cell}")
        c.number_format = CURRENCY_FMT
        c.font = Font(name="Arial", size=10, bold=True, color="000000")


# ── Monthly Allocation Sheet (managed services) ──────────────────────────────
def build_monthly_allocation_sheet(wb, config, rates, layout: Layout):
    """Managed services. Per-cell costs become formulas: daysPerMonth * dailyRate * uplift.
    Daily Rate column is a VLOOKUP into Config rate card.
    """
    ws = wb.create_sheet("Monthly Allocation")
    ws.sheet_properties.tabColor = NAVY

    allocs = config.get("monthlyAllocations", [])
    months = config.get("contractMonths", 12)
    start = datetime.strptime(config["startDate"], "%Y-%m-%d")
    annual_uplift = config.get("annualUplift", 0)

    headers = ["Function", "Rate Band", "Daily Rate", "Days / Month"]
    for m in range(months):
        d = add_month(start, m)
        headers.append(d.strftime("%b %Y"))
    headers += ["Total Days", "Total Cost"]

    widths = [24, 14, 14, 14] + [12] * months + [12, 16]
    set_col_widths(ws, widths)
    for i, h in enumerate(headers):
        ws.cell(row=1, column=i + 1, value=h)
    style_header_row(ws, 1, len(headers))

    c_func, c_rb, c_dr, c_dpm = 1, 2, 3, 4
    first_month_col = 5
    last_month_col = 4 + months
    c_total_days = 5 + months
    c_total_cost = 6 + months

    r = 2
    first_alloc_row = r
    for alloc_item in allocs:
        fn = alloc_item["function"]
        rb = alloc_item.get("rateBand", "Senior")
        dpm = alloc_item["daysPerMonth"]

        write_cell(ws, r, c_func, fn)
        write_cell(ws, r, c_rb, rb)
        # Daily Rate — VLOOKUP into rate card
        write_cell(ws, r, c_dr,
                   f"=VLOOKUP({col_letter(c_rb)}{r},{layout.rate_card_range},2,FALSE)",
                   fmt=CURRENCY_FMT)
        write_cell(ws, r, c_dpm, dpm, fmt=DEC_FMT)

        # Per-month cells = daysPerMonth * dailyRate * (1 + uplift)^year
        for m_idx in range(months):
            year_idx = m_idx // 12
            uplift_factor = (1 + annual_uplift) ** year_idx
            col = first_month_col + m_idx
            # daysPerMonth and dailyRate are cells; uplift is a hard-coded multiplier
            # (year-based, so safe; if you want yearly uplift to be configurable too,
            #  the Annual Uplift cell already exists in Config and could be referenced)
            formula = f"={col_letter(c_dpm)}{r}*{col_letter(c_dr)}{r}"
            if uplift_factor != 1.0:
                formula += f"*{uplift_factor}"
            write_cell(ws, r, col, formula, fmt=CURRENCY_FMT)

        # Total Days = days/month * months
        write_cell(ws, r, c_total_days,
                   f"={col_letter(c_dpm)}{r}*{months}",
                   fmt=DEC_FMT, font=bold_font)
        # Total Cost = SUM(month range)
        write_cell(ws, r, c_total_cost,
                   f"=SUM({col_letter(first_month_col)}{r}:{col_letter(last_month_col)}{r})",
                   fmt=CURRENCY_FMT, font=bold_font)
        r += 1
    last_alloc_row = r - 1

    # Grand totals
    style_total_row(ws, r, len(headers))
    ws.cell(row=r, column=1, value="TOTAL").font = teal_font
    for m_idx in range(months):
        col = first_month_col + m_idx
        L = col_letter(col)
        write_cell(ws, r, col,
                   f"=SUM({L}{first_alloc_row}:{L}{last_alloc_row})",
                   fmt=CURRENCY_FMT, font=teal_font, fill=teal_fill)
    L_td = col_letter(c_total_days)
    L_tc = col_letter(c_total_cost)
    write_cell(ws, r, c_total_days,
               f"=SUM({L_td}{first_alloc_row}:{L_td}{last_alloc_row})",
               fmt=DEC_FMT, font=teal_font, fill=teal_fill)
    write_cell(ws, r, c_total_cost,
               f"=SUM({L_tc}{first_alloc_row}:{L_tc}{last_alloc_row})",
               fmt=CURRENCY_FMT, font=teal_font, fill=teal_fill)
    grand_cost_cell = f"Monthly Allocation!${L_tc}${r}"  # before we change r
    total_row_ms = r

    r += 1
    ws.cell(row=r, column=c_total_days - 1, value="GST").font = bold_font
    write_cell(ws, r, c_total_cost,
               f"={L_tc}{total_row_ms}*{layout.gst_cell}",
               fmt=CURRENCY_FMT)
    r += 1
    ws.cell(row=r, column=c_total_days - 1, value="Total Inc. GST").font = hdr_font
    style_total_row(ws, r, len(headers), fill=navy_fill, font=hdr_font)
    write_cell(ws, r, c_total_cost,
               f"={L_tc}{total_row_ms}*(1+{layout.gst_cell})",
               fmt=CURRENCY_FMT, font=hdr_font, fill=navy_fill)

    # Annual summary
    r += 2
    ws.cell(row=r, column=1, value="Annual Summary").font = subtitle_font
    r += 1
    ann_headers = ["Year", "Monthly Cost", "Annual Cost", "Annual Inc GST", "Uplift Applied"]
    for i, h in enumerate(ann_headers):
        ws.cell(row=r, column=i + 1, value=h)
    style_header_row(ws, r, len(ann_headers))
    r += 1

    years = math.ceil(months / 12)
    gst = config.get("gst", 0.10)
    for y in range(years):
        rate_mult = (1 + annual_uplift) ** y
        # monthly cost = SUM of (dpm * rate * uplift) for that year
        # Use SUMPRODUCT against the alloc rows
        dpm_range = f"{col_letter(c_dpm)}{first_alloc_row}:{col_letter(c_dpm)}{last_alloc_row}"
        dr_range = f"{col_letter(c_dr)}{first_alloc_row}:{col_letter(c_dr)}{last_alloc_row}"
        monthly_formula = f"=SUMPRODUCT({dpm_range},{dr_range})"
        if rate_mult != 1.0:
            monthly_formula += f"*{rate_mult}"
        months_in_year = min(12, months - y * 12)
        annual_formula = f"=B{r}*{months_in_year}"
        annual_inc_gst_formula = f"=C{r}*(1+{layout.gst_cell})"

        write_cell(ws, r, 1, f"Year {y + 1}", font=bold_font)
        write_cell(ws, r, 2, monthly_formula, fmt=CURRENCY_FMT)
        write_cell(ws, r, 3, annual_formula, fmt=CURRENCY_FMT)
        write_cell(ws, r, 4, annual_inc_gst_formula, fmt=CURRENCY_FMT)
        write_cell(ws, r, 5, f"{annual_uplift*100:.0f}%" if y > 0 else "Base")
        r += 1


# ── Resource Summary ──────────────────────────────────────────────────────────
def build_resource_summary(wb, config, rates, layout: Layout):
    model = config["model"]
    if model == "managed_services":
        return  # Monthly Allocation sheet serves this purpose

    ws = wb.create_sheet("Resource Summary")
    ws.sheet_properties.tabColor = NAVY
    show_cont = model in ("sprint", "phase")

    ws["A1"] = "Summary by Function"
    ws["A1"].font = subtitle_font

    if show_cont:
        headers = ["Function", "Rate Band", "Daily Rate", "Total Days", "Estimate", "Estimate Inc. Contingency", "% of Total"]
        widths = [24, 14, 14, 12, 16, 22, 12]
    else:
        headers = ["Function", "Rate Band", "Daily Rate", "Total Days", "Estimate", "% of Total"]
        widths = [24, 14, 14, 12, 16, 12]

    set_col_widths(ws, widths)
    for i, h in enumerate(headers):
        ws.cell(row=2, column=i + 1, value=h)
    style_header_row(ws, 2, len(headers))

    # Group tasks by function (use first rate band seen)
    func_data = {}
    for t in config["tasks"]:
        fn = t["function"]
        rb = t.get("rateBand", "Senior")
        days = get_task_days(t, config)
        rate = rates.get(rb, 1963)
        est = days * rate
        cont = get_contingency(t.get("solutionArea", ""), config)
        est_c = est * (1 + cont)
        if fn not in func_data:
            func_data[fn] = {"rateBand": rb, "rate": rate, "days": 0, "est": 0, "est_c": 0}
        func_data[fn]["days"] += days
        func_data[fn]["est"] += est
        func_data[fn]["est_c"] += est_c

    # Sort once and capture order for formulas
    sorted_funcs = sorted(func_data.items(), key=lambda x: -x[1]["est"])

    # Refs into Estimate sheet
    est_sheet = "Estimate"
    days_L = col_letter(layout.est_days_col)
    est_L = col_letter(layout.est_estimate_col)
    func_L = col_letter(layout.est_func_col)
    rb_L = col_letter(layout.est_rateband_col)
    fr = layout.est_first_row
    lr = layout.est_last_row
    days_range = f"{est_sheet}!${days_L}${fr}:${days_L}${lr}"
    est_range = f"{est_sheet}!${est_L}${fr}:${est_L}${lr}"
    func_range = f"{est_sheet}!${func_L}${fr}:${func_L}${lr}"
    rb_range = f"{est_sheet}!${rb_L}${fr}:${rb_L}${lr}"
    if show_cont:
        ec_L = col_letter(layout.est_est_cont_col)
        ec_range = f"{est_sheet}!${ec_L}${fr}:${ec_L}${lr}"
        cascade_total_cell = f"{est_sheet}!${ec_L}${layout.est_total_row}"
    else:
        ec_range = ""
        cascade_total_cell = f"{est_sheet}!${est_L}${layout.est_total_row}"

    r = 3
    first_summary_row = r
    for fn, d in sorted_funcs:
        write_cell(ws, r, 1, fn)
        write_cell(ws, r, 2, d["rateBand"])
        # Daily Rate looked up from rate card
        write_cell(ws, r, 3,
                   f"=VLOOKUP(B{r},{layout.rate_card_range},2,FALSE)",
                   fmt=CURRENCY_FMT)
        # Total Days = SUMIFS(Estimate days, Estimate function = A{r})
        write_cell(ws, r, 4,
                   f"=SUMIFS({days_range},{func_range},A{r})",
                   fmt=DEC_FMT)
        # Estimate = SUMIFS
        write_cell(ws, r, 5,
                   f"=SUMIFS({est_range},{func_range},A{r})",
                   fmt=CURRENCY_FMT)
        if show_cont:
            write_cell(ws, r, 6,
                       f"=SUMIFS({ec_range},{func_range},A{r})",
                       fmt=CURRENCY_FMT)
            # % of total
            write_cell(ws, r, 7,
                       f"=IFERROR(F{r}/{cascade_total_cell},0)",
                       fmt=PCT_FMT)
        else:
            write_cell(ws, r, 6,
                       f"=IFERROR(E{r}/{cascade_total_cell},0)",
                       fmt=PCT_FMT)
        r += 1
    last_summary_row = r - 1

    # Totals
    style_total_row(ws, r, len(headers))
    ws.cell(row=r, column=1, value="TOTAL").font = teal_font
    write_cell(ws, r, 4,
               f"=SUM(D{first_summary_row}:D{last_summary_row})",
               fmt=DEC_FMT, font=teal_font, fill=teal_fill)
    write_cell(ws, r, 5,
               f"=SUM(E{first_summary_row}:E{last_summary_row})",
               fmt=CURRENCY_FMT, font=teal_font, fill=teal_fill)
    if show_cont:
        write_cell(ws, r, 6,
                   f"=SUM(F{first_summary_row}:F{last_summary_row})",
                   fmt=CURRENCY_FMT, font=teal_font, fill=teal_fill)
        write_cell(ws, r, 7,
                   f"=SUM(G{first_summary_row}:G{last_summary_row})",
                   fmt=PCT_FMT, font=teal_font, fill=teal_fill)
    else:
        write_cell(ws, r, 6,
                   f"=SUM(F{first_summary_row}:F{last_summary_row})",
                   fmt=PCT_FMT, font=teal_font, fill=teal_fill)

    # By Solution Area (fixed price only)
    if show_cont:
        sa_L = col_letter(layout.est_area_col)
        sa_range = f"{est_sheet}!${sa_L}${fr}:${sa_L}${lr}"
        cont_L = col_letter(layout.est_cont_pct_col)

        # Aggregate distinct solution areas in Python (for row labels)
        area_data = {}
        for t in config["tasks"]:
            sa = t.get("solutionArea", "Other")
            days = get_task_days(t, config)
            rate = rates.get(t.get("rateBand", "Senior"), 1963)
            est = days * rate
            cont = get_contingency(sa, config)
            est_c = est * (1 + cont)
            if sa not in area_data:
                area_data[sa] = {"cont": cont, "days": 0, "est": 0, "est_c": 0}
            area_data[sa]["days"] += days
            area_data[sa]["est"] += est
            area_data[sa]["est_c"] += est_c

        sa_start = r + 3
        ws.cell(row=sa_start, column=1, value="Summary by Solution Area").font = subtitle_font
        headers2 = ["Solution Area", "Contingency %", "Total Days", "Estimate", "Estimate Inc. Contingency", "% of Total"]
        for i, h in enumerate(headers2):
            ws.cell(row=sa_start + 1, column=i + 1, value=h)
        style_header_row(ws, sa_start + 1, len(headers2))

        r = sa_start + 2
        first_sa_row = r
        for sa, d in sorted(area_data.items(), key=lambda x: -x[1]["est_c"]):
            write_cell(ws, r, 1, sa)
            # Contingency % — keep as value to match what's on the row (could VLOOKUP but
            # this is illustrative; the source of truth lives in Estimate column)
            write_cell(ws, r, 2, d["cont"], fmt=PCT_FMT)
            write_cell(ws, r, 3,
                       f"=SUMIFS({days_range},{sa_range},A{r})",
                       fmt=DEC_FMT)
            write_cell(ws, r, 4,
                       f"=SUMIFS({est_range},{sa_range},A{r})",
                       fmt=CURRENCY_FMT)
            write_cell(ws, r, 5,
                       f"=SUMIFS({ec_range},{sa_range},A{r})",
                       fmt=CURRENCY_FMT)
            write_cell(ws, r, 6,
                       f"=IFERROR(E{r}/{cascade_total_cell},0)",
                       fmt=PCT_FMT)
            r += 1
        last_sa_row = r - 1

        style_total_row(ws, r, len(headers2))
        ws.cell(row=r, column=1, value="TOTAL").font = teal_font
        write_cell(ws, r, 3,
                   f"=SUM(C{first_sa_row}:C{last_sa_row})",
                   fmt=DEC_FMT, font=teal_font, fill=teal_fill)
        write_cell(ws, r, 4,
                   f"=SUM(D{first_sa_row}:D{last_sa_row})",
                   fmt=CURRENCY_FMT, font=teal_font, fill=teal_fill)
        write_cell(ws, r, 5,
                   f"=SUM(E{first_sa_row}:E{last_sa_row})",
                   fmt=CURRENCY_FMT, font=teal_font, fill=teal_fill)
        write_cell(ws, r, 6,
                   f"=SUM(F{first_sa_row}:F{last_sa_row})",
                   fmt=PCT_FMT, font=teal_font, fill=teal_fill)


# ── Weekly View ───────────────────────────────────────────────────────────────
def build_weekly_view(wb, config, rates, layout: Layout):
    ws = wb.create_sheet("Weekly View")
    ws.sheet_properties.tabColor = TEAL

    weeks, functions, weekly, weekly_cost, func_rateband = build_weekly_data(config, rates)

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 20
    for i in range(len(functions)):
        ws.column_dimensions[get_column_letter(3 + i)].width = 18

    total_col = 3 + len(functions)
    cost_col = total_col + 1
    ws.column_dimensions[get_column_letter(total_col)].width = 12
    ws.column_dimensions[get_column_letter(cost_col)].width = 16

    model_label = MODEL_LABELS.get(config["model"], "")
    ws["A1"] = f"Weekly Resource View — {model_label}"
    ws["A1"].font = title_font
    ws.cell(row=2, column=1, value=("Per-cell allocations are values (derived from working-day math). "
                                    "Row totals + grand totals are formulas. "
                                    "Colour: Green < 2.5d, Amber 2.5–4d, Red > 4d")).font = note_font

    # Row 3 = hidden RATE ROW. Holds per-function daily rate via VLOOKUP so weekly cost
    # uses the formula =SUMPRODUCT(allocs, rates). When you edit a daily rate in Config,
    # the weekly cost row recalculates.
    rate_row = 3
    ws.cell(row=rate_row, column=1, value="Daily Rates →").font = note_font
    for fi, fn in enumerate(functions):
        rb = func_rateband.get(fn, "Senior")
        # Two-step: write rate band to col B (visible label) — actually we use col B for "Phase"
        # so just write the formula referencing the rate band literal via VLOOKUP.
        # We embed the rate band as a string literal inside VLOOKUP to keep it self-contained.
        col = 3 + fi
        ws.cell(row=rate_row, column=col,
                value=f'=VLOOKUP("{rb}",{layout.rate_card_range},2,FALSE)').number_format = CURRENCY_FMT
    # Hide the rate row
    ws.row_dimensions[rate_row].hidden = True

    # Headers at row 4
    row = 4
    headers = ["Week Commencing", "Phase"] + functions + ["Total Days", "Total Cost"]
    for i, h in enumerate(headers):
        ws.cell(row=row, column=i + 1, value=h)
    style_header_row(ws, row, len(headers))

    row = 5
    first_data_row = row
    first_fn_col = 3
    last_fn_col = 2 + len(functions)

    for wk in weeks:
        phase = get_phase_for_week(config, wk)
        ws.cell(row=row, column=1, value=wk.strftime("%d %b %Y")).font = body_font
        ws.cell(row=row, column=1).border = thin_border
        ws.cell(row=row, column=2, value=phase).font = body_font
        ws.cell(row=row, column=2).border = thin_border

        for fi, fn in enumerate(functions):
            val = round(weekly[wk][fn], 1)
            cell = ws.cell(row=row, column=first_fn_col + fi, value=val if val > 0 else None)
            cell.font = body_font
            cell.border = thin_border
            cell.number_format = DEC_FMT
            if val > 4:
                cell.fill = red_fill
            elif val > 2.5:
                cell.fill = amber_fill
            elif val > 0:
                cell.fill = green_fill

        # Row total days = SUM
        first_fn_L = col_letter(first_fn_col)
        last_fn_L = col_letter(last_fn_col)
        write_cell(ws, row, total_col,
                   f"=SUM({first_fn_L}{row}:{last_fn_L}{row})",
                   fmt=DEC_FMT, font=bold_font)
        # Row total cost = SUMPRODUCT(allocs, rates)
        write_cell(ws, row, cost_col,
                   f"=SUMPRODUCT({first_fn_L}{row}:{last_fn_L}{row},{first_fn_L}{rate_row}:{last_fn_L}{rate_row})",
                   fmt=CURRENCY_FMT)
        row += 1

    last_data_row = row - 1

    # Grand total row
    style_total_row(ws, row, len(headers))
    ws.cell(row=row, column=1, value="TOTAL").font = teal_font
    for fi in range(len(functions)):
        col = first_fn_col + fi
        L = col_letter(col)
        write_cell(ws, row, col,
                   f"=SUM({L}{first_data_row}:{L}{last_data_row})",
                   fmt=DEC_FMT, font=teal_font, fill=teal_fill)
    total_L = col_letter(total_col)
    cost_L = col_letter(cost_col)
    write_cell(ws, row, total_col,
               f"=SUM({total_L}{first_data_row}:{total_L}{last_data_row})",
               fmt=DEC_FMT, font=teal_font, fill=teal_fill)
    write_cell(ws, row, cost_col,
               f"=SUM({cost_L}{first_data_row}:{cost_L}{last_data_row})",
               fmt=CURRENCY_FMT, font=teal_font, fill=teal_fill)
    ws.freeze_panes = "C5"


# ── Forecast Export ───────────────────────────────────────────────────────────
def build_forecast_export(wb, config, rates, layout: Layout):
    ws = wb.create_sheet("Forecast Export")
    ws.sheet_properties.tabColor = TEAL

    weeks, functions, weekly, weekly_cost, _ = build_weekly_data(config, rates)

    headers = ["Project", "Client", "Version", "Model", "Function", "Rate Band", "Week Commencing", "Estimated Days", "Daily Rate", "Estimated Cost"]
    widths = [22, 22, 10, 20, 24, 14, 16, 14, 14, 16]
    set_col_widths(ws, widths)
    for i, h in enumerate(headers):
        ws.cell(row=1, column=i + 1, value=h)
    style_header_row(ws, 1, len(headers))

    model = config["model"]
    func_meta = {}
    if model == "managed_services":
        for a in config.get("monthlyAllocations", []):
            func_meta[a["function"]] = {"rateBand": a.get("rateBand", "Senior")}
    else:
        for t in config.get("tasks", []):
            if t["function"] not in func_meta:
                func_meta[t["function"]] = {"rateBand": t.get("rateBand", "Senior")}

    row = 2
    proj = config.get("projectName", "")
    client = config.get("clientName", "")
    ver = config.get("estimateVersion", "V1")
    model_label = MODEL_LABELS.get(model, model)

    for wk in weeks:
        for fn in functions:
            days = round(weekly[wk][fn], 2)
            if days <= 0:
                continue
            meta = func_meta.get(fn, {})
            rb = meta.get("rateBand", "Senior")
            vals = [proj, client, ver, model_label, fn, rb,
                    wk.strftime("%Y-%m-%d"), round(days, 1)]
            write_row(ws, row, vals, fmt_map={8: DEC_FMT})
            # Daily Rate via VLOOKUP on Rate Band cell (col 6 = F)
            write_cell(ws, row, 9,
                       f"=VLOOKUP(F{row},{layout.rate_card_range},2,FALSE)",
                       fmt=CURRENCY_FMT)
            # Estimated Cost = days * rate
            write_cell(ws, row, 10,
                       f"=H{row}*I{row}",
                       fmt=CURRENCY_FMT)
            row += 1


# ── Payment Milestones (fixed price only) ─────────────────────────────────────
def build_milestones_sheet(wb, config, layout: Layout):
    milestones = config.get("paymentMilestones", [])
    if not milestones:
        return

    ws = wb.create_sheet("Payment Milestones")
    ws.sheet_properties.tabColor = NAVY

    headers = ["#", "Milestone", "% of Total", "Amount Ex GST", "Amount Inc GST", "Due Date", "Acceptance Criteria"]
    widths = [6, 28, 12, 18, 18, 14, 40]
    set_col_widths(ws, widths)
    for i, h in enumerate(headers):
        ws.cell(row=1, column=i + 1, value=h)
    style_header_row(ws, 1, len(headers))

    fee_ex = layout.est_fee_basis_ex_gst_cell  # e.g. "Estimate!$L$22"
    gst_cell = layout.gst_cell

    r = 2
    first_ms_row = r
    for i, m in enumerate(milestones):
        pct = m.get("percentage", 0)
        # 1, name, pct, amount ex gst (formula), amount inc gst (formula), due date, criteria
        write_cell(ws, r, 1, i + 1)
        write_cell(ws, r, 2, m.get("name", ""))
        write_cell(ws, r, 3, pct, fmt=PCT_FMT)
        # Amount Ex GST = fee_ex * pct
        write_cell(ws, r, 4, f"={fee_ex}*C{r}", fmt=CURRENCY_FMT)
        # Amount Inc GST = ex * (1 + GST%)
        write_cell(ws, r, 5, f"=D{r}*(1+{gst_cell})", fmt=CURRENCY_FMT)
        write_cell(ws, r, 6, m.get("dueDate", ""))
        write_cell(ws, r, 7, m.get("criteria", ""))
        r += 1
    last_ms_row = r - 1

    style_total_row(ws, r, len(headers))
    ws.cell(row=r, column=2, value="TOTAL").font = teal_font
    write_cell(ws, r, 3,
               f"=SUM(C{first_ms_row}:C{last_ms_row})",
               fmt=PCT_FMT, font=teal_font, fill=teal_fill)
    write_cell(ws, r, 4,
               f"=SUM(D{first_ms_row}:D{last_ms_row})",
               fmt=CURRENCY_FMT, font=teal_font, fill=teal_fill)
    write_cell(ws, r, 5,
               f"=SUM(E{first_ms_row}:E{last_ms_row})",
               fmt=CURRENCY_FMT, font=teal_font, fill=teal_fill)


# ── Main ──────────────────────────────────────────────────────────────────────
def calc_totals(config, rates):
    """Python-side totals — used ONLY for the printed summary at the end of main().
    Cells in the workbook use formulas; this function exists so the CLI output is
    helpful without opening Excel."""
    model = config["model"]
    gst = config.get("gst", 0.10)

    if model == "managed_services":
        allocs = config.get("monthlyAllocations", [])
        months = config.get("contractMonths", 12)
        annual_uplift = config.get("annualUplift", 0)
        total_days = sum(a["daysPerMonth"] for a in allocs) * months
        total_est = 0
        for a in allocs:
            rb = a.get("rateBand", "Senior")
            base_rate = rates.get(rb, 1963)
            for m in range(months):
                rate = base_rate * ((1 + annual_uplift) ** (m // 12))
                total_est += a["daysPerMonth"] * rate
        total_cont = total_est
    else:
        total_days = sum(get_task_days(t, config) for t in config["tasks"])
        total_est = sum(get_task_days(t, config) * rates.get(t.get("rateBand", "Senior"), 1963) for t in config["tasks"])
        total_cont = sum(
            get_task_days(t, config) * rates.get(t.get("rateBand", "Senior"), 1963)
            * (1 + get_contingency(t.get("solutionArea", ""), config))
            for t in config["tasks"]
        )
        discount = get_discount(config)
        if discount > 0:
            total_est *= (1 - discount)
            total_cont *= (1 - discount)

    return total_days, total_est, total_cont


def main():
    if len(sys.argv) < 3:
        print("Usage: python build-estimate.py <config.json> <output.xlsx>")
        sys.exit(1)

    try:
        with open(sys.argv[1]) as f:
            config = json.load(f)
    except (json.JSONDecodeError, FileNotFoundError) as e:
        print(f"Error reading config: {e}")
        sys.exit(1)

    model = config.get("model", "sprint")
    if model not in ("sprint", "phase", "t&m", "managed_services"):
        print(f"Unknown model: {model}. Expected: sprint, phase, t&m, managed_services")
        sys.exit(1)

    fy = config.get("financialYear", "FY26")
    rates = RATE_CARDS.get(fy, RATE_CARDS["FY26"])

    layout = Layout()

    wb = Workbook()
    build_config_sheet(wb, config, rates, layout)
    build_estimate_sheet(wb, config, rates, layout)
    build_resource_summary(wb, config, rates, layout)
    build_weekly_view(wb, config, rates, layout)
    build_forecast_export(wb, config, rates, layout)

    total_days, total_est, total_cont = calc_totals(config, rates)
    gst = config.get("gst", 0.10)

    if model in ("sprint", "phase"):
        build_milestones_sheet(wb, config, layout)

    wb.save(sys.argv[2])
    print(f"Workbook saved: {sys.argv[2]}")
    print(f"Model: {MODEL_LABELS.get(model, model)}")
    discount = get_discount(config)
    print(f"\nSummary:")
    print(f"  Total Days: {total_days:.0f}")
    print(f"  Estimate (ex contingency): ${total_est:,.0f}")
    if discount > 0:
        print(f"  Discount: {discount*100:.0f}% applied")
    if model in ("sprint", "phase"):
        print(f"  Estimate (inc contingency): ${total_cont:,.0f}")
        print(f"  GST: ${total_cont * gst:,.0f}")
        print(f"  Total Inc GST: ${total_cont * (1 + gst):,.0f}")
    elif model == "t&m":
        print(f"  GST: ${total_est * gst:,.0f}")
        print(f"  Total Inc GST: ${total_est * (1 + gst):,.0f}")
        if config.get("budgetCap"):
            cap = config["budgetCap"]
            cap_ex = cap / (1 + gst) if config.get("budgetCapInclGST") else cap
            print(f"  Budget Cap (Ex GST): ${cap_ex:,.0f}")
            print(f"  Remaining: ${cap_ex - total_est:,.0f}")
    elif model == "managed_services":
        print(f"  Monthly Cost: ${total_est / config.get('contractMonths', 12):,.0f}")
        print(f"  GST: ${total_est * gst:,.0f}")
        print(f"  Total Inc GST: ${total_est * (1 + gst):,.0f}")


if __name__ == "__main__":
    main()
