#!/usr/bin/env python3
"""Add a visual project timeline (Gantt-style) sheet to an existing estimate workbook."""

import json, sys
from datetime import datetime, timedelta
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

NAVY = "19263C"
TEAL = "6BE1B8"
WHITE = "FFFFFF"
LIGHT_GREY = "F2F2F2"

navy_fill = PatternFill("solid", fgColor=NAVY)
teal_fill = PatternFill("solid", fgColor=TEAL)
grey_fill = PatternFill("solid", fgColor=LIGHT_GREY)

phase_colors = [
    PatternFill("solid", fgColor="4472C4"),  # Blue - Discovery
    PatternFill("solid", fgColor="ED7D31"),  # Orange - Build
    PatternFill("solid", fgColor="A5A5A5"),  # Grey - UAT
    PatternFill("solid", fgColor="70AD47"),  # Green - Go-Live
]
milestone_fill = PatternFill("solid", fgColor="FFD700")

title_font = Font(name="Arial", size=14, bold=True, color=NAVY)
subtitle_font = Font(name="Arial", size=11, bold=True, color=NAVY)
hdr_font = Font(name="Arial", size=9, bold=True, color=WHITE)
body_font = Font(name="Arial", size=9)
bold_font = Font(name="Arial", size=9, bold=True)
phase_font = Font(name="Arial", size=8, bold=True, color=WHITE)
milestone_font = Font(name="Arial", size=8, bold=True, color="000000")

thin_border = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)


def add_timeline(wb, config):
    ws = wb.create_sheet("Project Timeline", 0)  # Insert as first sheet
    ws.sheet_properties.tabColor = NAVY

    phases = config.get("phases", [])
    milestones = config.get("paymentMilestones", [])
    project = config.get("projectName", "Project")
    client = config.get("clientName", "Client")

    if not phases:
        return

    # Calculate date range
    all_starts = [datetime.strptime(p["startDate"], "%Y-%m-%d") for p in phases]
    all_ends = [datetime.strptime(p["endDate"], "%Y-%m-%d") for p in phases]
    proj_start = min(all_starts)
    proj_end = max(all_ends)

    # Align to week boundaries (Monday)
    chart_start = proj_start - timedelta(days=proj_start.weekday())
    chart_end = proj_end + timedelta(days=(6 - proj_end.weekday()))

    total_weeks = ((chart_end - chart_start).days // 7) + 1
    week_col_start = 4  # Columns A-C for labels, D onwards for weeks

    # Column widths
    from openpyxl.utils import get_column_letter
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 14
    for w in range(total_weeks):
        col_letter = get_column_letter(week_col_start + w)
        ws.column_dimensions[col_letter].width = 12

    # Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(week_col_start + total_weeks - 1, 16))
    ws["A1"] = f"{project} — Project Timeline"
    ws["A1"].font = title_font

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=min(week_col_start + total_weeks - 1, 16))
    ws["A2"] = f"{client} | {phases[0]['startDate']} to {phases[-1]['endDate']}"
    ws["A2"].font = Font(name="Arial", size=10, color="666666")

    # Week headers
    header_row = 4
    ws.cell(row=header_row, column=1, value="#").font = hdr_font
    ws.cell(row=header_row, column=2, value="Item").font = hdr_font
    ws.cell(row=header_row, column=3, value="Duration").font = hdr_font
    for c in range(1, 4):
        ws.cell(row=header_row, column=c).fill = navy_fill
        ws.cell(row=header_row, column=c).font = hdr_font
        ws.cell(row=header_row, column=c).border = thin_border
        ws.cell(row=header_row, column=c).alignment = Alignment(horizontal="center", vertical="center")

    for w in range(total_weeks):
        col = week_col_start + w
        week_date = chart_start + timedelta(weeks=w)
        cell = ws.cell(row=header_row, column=col, value=week_date.strftime("%d %b"))
        cell.fill = navy_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Month header row (row 3)
    month_row = 3
    current_month = None
    month_start_col = None
    for w in range(total_weeks):
        col = week_col_start + w
        week_date = chart_start + timedelta(weeks=w)
        month_label = week_date.strftime("%B %Y")
        if month_label != current_month:
            if current_month is not None and month_start_col is not None:
                if col - 1 > month_start_col:
                    ws.merge_cells(start_row=month_row, start_column=month_start_col, end_row=month_row, end_column=col - 1)
            current_month = month_label
            month_start_col = col
            cell = ws.cell(row=month_row, column=col, value=month_label)
            cell.font = subtitle_font
            cell.alignment = Alignment(horizontal="center")
    # Close last month
    if month_start_col is not None:
        last_col = week_col_start + total_weeks - 1
        if last_col > month_start_col:
            ws.merge_cells(start_row=month_row, start_column=month_start_col, end_row=month_row, end_column=last_col)

    # Phase bars
    row = header_row + 1
    for pi, phase in enumerate(phases):
        p_start = datetime.strptime(phase["startDate"], "%Y-%m-%d")
        p_end = datetime.strptime(phase["endDate"], "%Y-%m-%d")
        duration_days = (p_end - p_start).days + 1
        weeks_count = max(1, round(duration_days / 7))

        ws.cell(row=row, column=1, value=pi + 1).font = bold_font
        ws.cell(row=row, column=1).border = thin_border
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=2, value=phase["name"]).font = bold_font
        ws.cell(row=row, column=2).border = thin_border
        ws.cell(row=row, column=3, value=f"{weeks_count} wks").font = body_font
        ws.cell(row=row, column=3).border = thin_border
        ws.cell(row=row, column=3).alignment = Alignment(horizontal="center")

        color = phase_colors[pi % len(phase_colors)]
        for w in range(total_weeks):
            col = week_col_start + w
            week_start = chart_start + timedelta(weeks=w)
            week_end = week_start + timedelta(days=6)
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border

            if week_start <= p_end and week_end >= p_start:
                cell.fill = color
                # Put phase name in first bar cell
                if week_start <= p_start <= week_end or (w == 0 and p_start < chart_start):
                    cell.value = phase["name"]
                    cell.font = phase_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.fill = PatternFill("solid", fgColor="FAFAFA")

        # Row height
        ws.row_dimensions[row].height = 28
        row += 1

    # Spacer
    row += 1

    # Milestone rows
    ws.cell(row=row, column=1).font = bold_font
    ws.cell(row=row, column=2, value="Payment Milestones").font = subtitle_font
    row += 1

    for mi, ms in enumerate(milestones):
        ws.cell(row=row, column=1, value=f"M{mi + 1}").font = bold_font
        ws.cell(row=row, column=1).border = thin_border
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")

        name = ms.get("name", "").replace(f"M{mi+1} — ", "").replace(f"M{mi+1} - ", "")
        pct = ms.get("percentage", 0)
        ws.cell(row=row, column=2, value=f"{name} ({int(pct*100)}%)").font = body_font
        ws.cell(row=row, column=2).border = thin_border
        ws.cell(row=row, column=3, value=ms.get("dueDate", "")).font = body_font
        ws.cell(row=row, column=3).border = thin_border
        ws.cell(row=row, column=3).alignment = Alignment(horizontal="center")

        due = ms.get("dueDate")
        if due:
            due_date = datetime.strptime(due, "%Y-%m-%d")
            for w in range(total_weeks):
                col = week_col_start + w
                week_start = chart_start + timedelta(weeks=w)
                week_end = week_start + timedelta(days=6)
                cell = ws.cell(row=row, column=col)
                cell.border = thin_border

                if week_start <= due_date <= week_end:
                    cell.fill = milestone_fill
                    cell.value = "*"
                    cell.font = milestone_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.fill = PatternFill("solid", fgColor="FAFAFA")

        ws.row_dimensions[row].height = 22
        row += 1

    # Legend
    row += 2
    ws.cell(row=row, column=2, value="Legend").font = subtitle_font
    row += 1
    for pi, phase in enumerate(phases):
        cell = ws.cell(row=row, column=2)
        cell.fill = phase_colors[pi % len(phase_colors)]
        cell.font = phase_font
        cell.value = f"  {phase['name']}  "
        ws.cell(row=row, column=3, value=f"{phase['startDate']} to {phase['endDate']}").font = body_font
        row += 1
    cell = ws.cell(row=row, column=2)
    cell.fill = milestone_fill
    cell.font = milestone_font
    cell.value = "  * Milestone  "
    ws.cell(row=row, column=3, value="Payment milestone due date").font = body_font

    # Today marker in header
    today = datetime.now()
    if chart_start <= today <= chart_end:
        for w in range(total_weeks):
            week_start = chart_start + timedelta(weeks=w)
            week_end = week_start + timedelta(days=6)
            if week_start <= today <= week_end:
                col = week_col_start + w
                cell = ws.cell(row=header_row, column=col)
                cell.value = f"v {week_start.strftime('%d %b')}"
                cell.fill = PatternFill("solid", fgColor=TEAL)
                cell.font = Font(name="Arial", size=9, bold=True, color="000000")
                break

    ws.freeze_panes = "D5"
    print(f"Timeline sheet added with {len(phases)} phases and {len(milestones)} milestones")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python add-timeline.py <config.json> <workbook.xlsx>")
        sys.exit(1)

    with open(sys.argv[1]) as f:
        config = json.load(f)

    wb = load_workbook(sys.argv[2])
    add_timeline(wb, config)
    wb.save(sys.argv[2])
    print(f"Saved: {sys.argv[2]}")
