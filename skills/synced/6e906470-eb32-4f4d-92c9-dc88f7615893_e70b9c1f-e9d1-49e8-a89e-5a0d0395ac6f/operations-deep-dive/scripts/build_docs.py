# -*- coding: utf-8 -*-
"""Build the operations deep dive Word document(s) from a content JSON file.

Usage:
    python build_docs.py <content.json>

Produces one .docx per entry in content["audiences"] ("confidential" and/or
"team"). The confidential version includes the annex and all plan actions; the
team version drops the annex, omits confidential-only actions, and substitutes
sanitised wording where a "team_action" is provided.
"""
import sys
import json
import datetime as dt
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import nsdecls, qn
from docx.oxml import parse_xml, OxmlElement

# ---- Brand (Accelerate 2026) ----
NAVY = RGBColor(0x1A, 0x2B, 0x46)
BLUE = RGBColor(0x1D, 0x69, 0xF3)
DEEP = RGBColor(0x24, 0x3A, 0x5E)
GREY = RGBColor(0x55, 0x5B, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
FONT = "Figtree"
NAVY_HEX = "1A2B46"
LTGREY_HEX = "EEF1F5"

MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]


def parse_date(s):
    """Parse '25 May 2026' (with optional year)."""
    s = s.strip()
    for fmt in ("%d %B %Y", "%d %b %Y", "%Y-%m-%d", "%d %B", "%d %b"):
        try:
            d = dt.datetime.strptime(s, fmt)
            if d.year == 1900:
                d = d.replace(year=dt.date.today().year)
            return d.date()
        except ValueError:
            continue
    raise ValueError(f"Could not parse date: {s!r}")


def fmt_date(d, with_year=False):
    base = f"{d.day} {MONTHS[d.month - 1]}"
    return f"{base} {d.year}" if with_year else base


def phase_meta(plan_start):
    """Return [(label, daterange), ...] for the three phases."""
    s = parse_date(plan_start)
    spans = [(0, 30), (31, 60), (61, 90)]
    labels = ["Days 0-30", "Days 31-60", "Days 61-90"]
    out = []
    for (a, b), lab in zip(spans, labels):
        d0 = s + dt.timedelta(days=a)
        d1 = s + dt.timedelta(days=b)
        wy = d1.year != d0.year
        out.append((lab, f"{fmt_date(d0)} - {fmt_date(d1, with_year=wy)}"))
    return out


# ---- helpers ----
def shade(cell, hex_fill):
    cell._element.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_fill}"/>'))


def set_cell_margins(cell, top=60, bottom=60, left=110, right=110):
    tcPr = cell._element.get_or_add_tcPr()
    m = OxmlElement('w:tcMar')
    for tag, val in (('top', top), ('bottom', bottom), ('start', left), ('end', right)):
        e = OxmlElement(f'w:{tag}')
        e.set(qn('w:w'), str(val)); e.set(qn('w:type'), 'dxa'); m.append(e)
    tcPr.append(m)


def style_run(run, size=10, color=BLACK, bold=False, italic=False):
    run.font.name = FONT; run.font.size = Pt(size); run.font.color.rgb = color
    run.bold = bold; run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.append(rfonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rfonts.set(qn(a), FONT)


def setup_styles(doc):
    normal = doc.styles['Normal']
    normal.font.name = FONT; normal.font.size = Pt(10); normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.12
    for lvl, sz in ((1, 18), (2, 14), (3, 12)):
        st = doc.styles[f'Heading {lvl}']
        st.font.name = FONT; st.font.size = Pt(sz); st.font.bold = True; st.font.color.rgb = NAVY
        st.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
        st.paragraph_format.space_after = Pt(4); st.paragraph_format.keep_with_next = True


def add_para(doc, text="", size=10, color=BLACK, bold=False, italic=False, space_after=6):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space_after)
    if text:
        style_run(p.add_run(text), size=size, color=color, bold=bold, italic=italic)
    return p


def add_bullet(doc, text, lead=None, size=10):
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(3)
    if lead:
        style_run(p.add_run(lead), size=size, color=NAVY, bold=True)
    style_run(p.add_run(text), size=size, color=BLACK)
    return p


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.name = FONT; r.font.color.rgb = NAVY
    return h


def add_rule(doc):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8); p.paragraph_format.space_before = Pt(2)
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), NAVY_HEX)
    pbdr.append(bottom); pPr.append(pbdr)


def add_table(doc, headers, rows, widths=None, body_size=9, emphasis_col=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.style = 'Table Grid'
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]; cell.text = ""; set_cell_margins(cell)
        style_run(cell.paragraphs[0].add_run(h), size=9, color=WHITE, bold=True)
        shade(cell, NAVY_HEX)
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci, val in enumerate(row):
            cell = cells[ci]; cell.text = ""; set_cell_margins(cell)
            emph = (emphasis_col is not None and ci == emphasis_col)
            style_run(cell.paragraphs[0].add_run(str(val)), size=body_size,
                      color=(NAVY if emph else BLACK), bold=emph)
            if ri % 2 == 1:
                shade(cell, LTGREY_HEX)
    if widths:
        for ci, w in enumerate(widths):
            for row in table.rows:
                row.cells[ci].width = Cm(w)
    return table


def add_cover(doc, cfg, confidential):
    style_run(add_para(doc, "ACCELERATE TECH", size=11, color=NAVY, bold=True, space_after=2).runs[0],
              size=11, color=NAVY, bold=True)
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]; cell.width = Cm(17); shade(cell, NAVY_HEX)
    set_cell_margins(cell, top=300, bottom=300, left=260, right=260); cell.text = ""
    style_run(cell.paragraphs[0].add_run(cfg.get("title", "Operations deep dive")), size=30, color=WHITE, bold=True)
    sp = cell.add_paragraph(); sp.paragraph_format.space_before = Pt(4)
    style_run(sp.add_run(cfg.get("subtitle", "Workshop summary and 30/60/90 plan")), size=15,
              color=RGBColor(0xCF, 0xDA, 0xEC))
    add_para(doc, "", space_after=6)
    aud = (f"Confidential briefing for {cfg.get('ceo_name','the CEO')} and {cfg.get('owner_name','the operations owner')}"
           if confidential else "Operations briefing for the leadership and delivery leads")
    if cfg.get("quarter_label"):
        aud = f"{cfg['quarter_label']}  -  {aud}"
    add_para(doc, aud, size=11, color=DEEP, bold=True, space_after=2)
    add_para(doc, f"Document date: {cfg.get('doc_date','')}", size=10, color=GREY, space_after=2)
    label = (f"Confidential - {cfg.get('ceo_name','CEO')} and {cfg.get('owner_name','owner')} only"
             if confidential else "For the operations owner, the practice and technical leads, and wider leadership")
    add_para(doc, label, size=9, color=(BLUE if confidential else GREY),
             bold=confidential, italic=not confidential, space_after=2)
    add_rule(doc)


def add_footer(doc, confidential):
    p = doc.sections[0].footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    txt = "Operations deep dive - Confidential" if confidential else "Operations deep dive"
    style_run(p.add_run(txt + "    |    Page "), size=8, color=GREY)
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
    r = OxmlElement('w:r'); r.append(OxmlElement('w:rPr')); te = OxmlElement('w:t'); te.text = "1"
    r.append(te); fld.append(r); p._p.append(fld)


def team_actions(actions):
    out = []
    for a in actions:
        if a.get("confidential"):
            if a.get("team_action"):
                out.append({**a, "action": a["team_action"]})
            # else: drop from team version
        else:
            out.append(a)
    return out


def build(content, confidential, out_path):
    cfg = content.get("config", {})
    doc = Document(); setup_styles(doc)
    section = doc.sections[0]
    section.page_width = Emu(int(21.0 * 360000)); section.page_height = Emu(int(29.7 * 360000))
    section.left_margin = Cm(2.0); section.right_margin = Cm(2.0)
    section.top_margin = Cm(1.8); section.bottom_margin = Cm(1.8)

    add_cover(doc, cfg, confidential); add_footer(doc, confidential)

    # 1. Purpose and context
    add_heading(doc, "1. Purpose and context", 1)
    ctx = cfg.get("workshop_context", "Recently")
    add_para(doc, f"{ctx}, {cfg.get('ceo_name','the CEO')} and {cfg.get('owner_name','the operations owner')} ran an operations deep dive: a focused war room to look at how delivery and operations run today, protect what is working, and fix what is not. This document is the record of that session. It sets out the wins worth keeping, the gaps costing us time and margin, the decisions taken in the room, and a 30 / 60 / 90 day plan with a clear owner and measure against every commitment.")
    metas = phase_meta(cfg["plan_start"])
    add_para(doc, f"The plan runs from {cfg['plan_start']}. The first 30 days close on {metas[0][1].split(' - ')[1]}, day 60 on {metas[1][1].split(' - ')[1]}, and day 90 on {metas[2][1].split(' - ')[1]}. The structural and people decisions are deliberately front-loaded so they are set early.")

    # 2. Wins
    add_heading(doc, "2. What's working well", 1)
    add_para(doc, "The session opened on what is going right, because plenty is.")
    for w in content.get("wins", []):
        add_bullet(doc, w)

    # 3. Themes
    add_heading(doc, "3. What needs to change", 1)
    themes = content.get("themes", [])
    add_para(doc, f"The rest of the session named the gaps. They group into {len(themes)} themes.")
    for i, th in enumerate(themes, start=1):
        add_heading(doc, f"3.{i} {th['title']}", 2)
        add_para(doc, th["body"])

    # 4. Decisions
    add_heading(doc, "4. Decisions made", 1)
    add_para(doc, "These decisions came out of the room. They are settled, not up for re-litigation.")
    for d in content.get("decisions", []):
        add_bullet(doc, d)

    # 5. KPIs
    add_heading(doc, "5. Operating KPIs", 1)
    add_para(doc, "These are the numbers we run operations against. Each one has a target, an owner, and a reporting cadence so there is no ambiguity about who is on the hook or how often we look.")
    kpis = content.get("kpis", [])
    add_table(doc, ["KPI", "Target", "Owner", "Cadence", "Source"],
              [[k["kpi"], k["target"], k["owner"], k["cadence"], k["source"]] for k in kpis],
              widths=[4.6, 3.6, 3.5, 2.6, 2.7], emphasis_col=0)

    # 6. Accountability
    acc = content.get("accountability", {})
    add_heading(doc, "6. Accountability model", 1)
    add_para(doc, "Accountability sits in three layers.")
    if acc.get("ceo"):
        add_heading(doc, acc["ceo"]["label"], 3); add_para(doc, acc["ceo"]["scope"])
    if acc.get("owner"):
        add_heading(doc, acc["owner"]["label"], 3); add_para(doc, acc["owner"]["scope"])
    if acc.get("team"):
        add_heading(doc, "The team", 3)
        for member in acc["team"]:
            add_bullet(doc, member["scope"], lead=member["role"] + ": ")

    # 7. Plan
    add_heading(doc, "7. The 30 / 60 / 90 plan", 1)
    for ph, (label, daterange) in zip(content.get("phases", []), metas):
        add_heading(doc, f"{label}: {ph['title']} ({daterange})", 2)
        if ph.get("intro"):
            add_para(doc, ph["intro"])
        acts = ph["actions"] if confidential else team_actions(ph["actions"])
        add_table(doc, ["Action", "Owner", "Done when"],
                  [[a["action"], a["owner"], a["done_when"]] for a in acts],
                  widths=[9.4, 3.8, 3.8], emphasis_col=1)

    # 8. Ideas
    if content.get("ideas"):
        add_heading(doc, "8. Ideas to pursue", 1)
        add_para(doc, "A few ideas surfaced that sit alongside the plan rather than inside it.")
        for idea in content["ideas"]:
            add_bullet(doc, idea)

    # Appendix - org chart
    img = cfg.get("org_chart_image")
    if img:
        import os
        doc.add_page_break()
        add_heading(doc, "Appendix: team structure", 1)
        add_para(doc, "Accelerate Tech organisational structure.")
        if os.path.exists(img):
            doc.add_picture(img, width=Cm(14.0))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            add_para(doc, f"[Org chart image not found: {img}]", italic=True, color=GREY)

    # Confidential annex
    if confidential and content.get("annex"):
        doc.add_page_break()
        add_heading(doc, "Confidential annex", 1)
        add_para(doc, f"This annex is for {cfg.get('ceo_name','the CEO')} and {cfg.get('owner_name','the operations owner')} only. It is not included in the version shared with the leads or wider leadership.",
                 italic=True, color=BLUE)
        for block in content["annex"]:
            add_heading(doc, block["heading"], 3)
            for item in block["items"]:
                add_bullet(doc, item)

    doc.save(out_path)
    print("Saved:", out_path)


def main():
    if len(sys.argv) < 2:
        print("Usage: python build_docs.py <content.json>"); sys.exit(1)
    with open(sys.argv[1], encoding="utf-8") as f:
        content = json.load(f)
    cfg = content.get("config", {})
    out_dir = cfg.get("output_dir", ".")
    stem = cfg.get("file_stem", "Operations Deep Dive - Workshop Summary and 30-60-90")
    import os
    for aud in content.get("audiences", ["confidential", "team"]):
        confidential = (aud.lower() == "confidential")
        out = os.path.join(out_dir, f"{stem} ({aud.capitalize()}).docx")
        build(content, confidential, out)


if __name__ == "__main__":
    main()
