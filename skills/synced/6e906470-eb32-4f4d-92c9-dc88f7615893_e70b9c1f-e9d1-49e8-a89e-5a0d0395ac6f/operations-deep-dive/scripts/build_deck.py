# -*- coding: utf-8 -*-
"""Build the operations deep dive readout deck (.pptx) from a content JSON file.

Usage:
    python build_deck.py <content.json>

Reads content["deck"] for the curated, short slide bullets (the deck is a
readout companion to the document - keep bullets terse) and content["config"]
for the cover. The deck is always team-safe; never put confidential content in it.

Builds on the Accelerate 2026 dark template. On Windows there is no Linux pptx
sandbox, so this drives template.pptx directly with python-pptx: it duplicates
the clean text-only example slides (cover, section divider) and uses the
'Title and Content' layout for body slides to avoid leftover example shapes.
"""
import sys
import os
import json
import copy
from pptx import Presentation
from pptx.util import Inches
from pptx.oxml.ns import qn
from lxml import etree

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))


def _find_template():
    """Locate the Accelerate 2026 deck template by walking up to the skills repo
    root, so this works regardless of how deeply this skill folder is nested."""
    rel = os.path.join("1-sales-marketing", "presentations",
                       "accelerate-slides", "assets", "template.pptx")
    d = SCRIPT_DIR
    while True:
        cand = os.path.join(d, rel)
        if os.path.exists(cand):
            return cand
        parent = os.path.dirname(d)
        if parent == d:
            return os.path.join(SCRIPT_DIR, rel)  # last-resort guess
        d = parent


DEFAULT_TEMPLATE = _find_template()

# template slide indices (clean, text-only examples) and layout indices
SRC_COVER = 2          # 'Cover - no photo'
SRC_DIVIDER = 14       # 'Section Divider'
LAYOUT_CONTENT = 13    # 'Title and Content' (title idx 0, body idx 11)


def duplicate_slide(prs, index):
    src = prs.slides[index]
    new = prs.slides.add_slide(src.slide_layout)
    for sp in list(new.shapes):
        sp._element.getparent().remove(sp._element)
    for sp in src.shapes:
        new.shapes._spTree.append(copy.deepcopy(sp._element))
    return new


def get_ph(slide, idx):
    for sh in slide.placeholders:
        if sh.placeholder_format.idx == idx:
            return sh
    raise KeyError(f"placeholder idx {idx} not found")


def set_ph_lines(slide, idx, lines):
    ph = get_ph(slide, idx)
    txBody = ph.text_frame._txBody
    paras = txBody.findall(qn('a:p'))
    template = copy.deepcopy(paras[0])
    for p in paras:
        txBody.remove(p)
    for line in lines:
        newp = copy.deepcopy(template)
        for br in newp.findall(qn('a:br')):
            newp.remove(br)
        runs = newp.findall(qn('a:r'))
        if runs:
            t = runs[0].find(qn('a:t'))
            if t is None:
                t = etree.SubElement(runs[0], qn('a:t'))
            t.text = line
            for extra in runs[1:]:
                newp.remove(extra)
        else:
            r = etree.SubElement(newp, qn('a:r'))
            etree.SubElement(r, qn('a:t')).text = line
        txBody.append(newp)


def set_title(slide, text):
    get_ph(slide, 0).text = text


def add_bullets(slide, items, size=18):
    """items = list of (text, level)."""
    ph = get_ph(slide, 11)
    txBody = ph.text_frame._txBody
    for p in txBody.findall(qn('a:p')):
        txBody.remove(p)
    for text, level in items:
        p = etree.SubElement(txBody, qn('a:p'))
        pPr = etree.SubElement(p, qn('a:pPr'))
        pPr.set('lvl', str(level))
        if level == 0:
            pPr.set('marL', '342900'); pPr.set('indent', '-342900'); bch = '•'
        else:
            pPr.set('marL', '742950'); pPr.set('indent', '-285750'); bch = '–'
        sb = etree.SubElement(pPr, qn('a:spcBef'))
        etree.SubElement(sb, qn('a:spcPts')).set('val', '600')
        bf = etree.SubElement(pPr, qn('a:buFont')); bf.set('typeface', 'Arial')
        etree.SubElement(pPr, qn('a:buChar')).set('char', bch)
        r = etree.SubElement(p, qn('a:r'))
        rPr = etree.SubElement(r, qn('a:rPr'))
        rPr.set('lang', 'en-AU')
        rPr.set('sz', str(int((size if level == 0 else size - 1) * 100)))
        rf = etree.SubElement(rPr, qn('a:rFonts')); rf.set('typeface', 'Figtree'); rf.set('cs', 'Figtree')
        fill = etree.SubElement(rPr, qn('a:solidFill'))
        etree.SubElement(fill, qn('a:srgbClr')).set('val', 'FFFFFF')
        etree.SubElement(r, qn('a:t')).text = text


def content_slide(prs, title, bullets, size=18):
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_CONTENT])
    set_title(s, title)
    add_bullets(s, [(b, 0) if isinstance(b, str) else (b["text"], b.get("level", 0))
                    for b in bullets], size=size)
    return s


def divider(prs, title):
    d = duplicate_slide(prs, SRC_DIVIDER)
    set_ph_lines(d, 13, [title])   # roomier upper box
    set_ph_lines(d, 12, [""])      # leave lower box empty (avoids overlap)
    return d


def chunk(seq, n):
    return [seq[i:i + n] for i in range(0, len(seq), n)]


def phase_titles(plan_start):
    import datetime as dt
    months = ["January", "February", "March", "April", "May", "June", "July",
              "August", "September", "October", "November", "December"]
    for fmt in ("%d %B %Y", "%d %b %Y", "%Y-%m-%d", "%d %B", "%d %b"):
        try:
            s = dt.datetime.strptime(plan_start.strip(), fmt).date(); break
        except ValueError:
            continue
    else:
        return ["Days 0-30", "Days 31-60", "Days 61-90"]
    labels = ["Days 0-30", "Days 31-60", "Days 61-90"]
    spans = [(0, 30), (31, 60), (61, 90)]
    out = []
    for (a, b), lab in zip(spans, labels):
        d0 = s + dt.timedelta(days=a); d1 = s + dt.timedelta(days=b)
        out.append(f"{lab} ({d0.day} {months[d0.month-1]} - {d1.day} {months[d1.month-1]})")
    return out


def build(content, out_path, template):
    cfg = content.get("config", {})
    deck = content.get("deck", {})
    prs = Presentation(template)
    n_orig = len(prs.slides._sldIdLst)

    # Cover
    cover = duplicate_slide(prs, SRC_COVER)
    set_ph_lines(cover, 0, [cfg.get("title", "Operations deep dive")])
    set_ph_lines(cover, 11, [cfg.get("subtitle", "Workshop summary and 30/60/90 plan")])
    _sub = get_ph(cover, 11)
    _sub.left = Inches(0.64); _sub.top = Inches(4.15); _sub.width = Inches(6.5)
    set_ph_lines(cover, 10, [f"{cfg.get('ceo_name','')} and {cfg.get('owner_name','')}".strip(" and"),
                             "Accelerate Tech", cfg.get("doc_date", "")])

    # Agenda
    content_slide(prs, "Agenda", [
        "What's working well",
        "What needs to change",
        "Decisions made in the room",
        "Operating KPIs",
        "The 30 / 60 / 90 plan and who owns what",
    ], size=20)

    # Wins
    if deck.get("wins"):
        content_slide(prs, "What's working well", deck["wins"], size=18)

    # Themes
    divider(prs, "What needs to change")
    if deck.get("themes"):
        content_slide(prs, "Six themes" if len(deck["themes"]) == 6 else "What needs to change",
                      deck["themes"], size=18)

    # Decisions
    if deck.get("decisions"):
        content_slide(prs, "Decisions made", deck["decisions"], size=18)

    # KPIs (chunk to keep slides readable)
    for i, grp in enumerate(chunk(deck.get("kpis", []), 8)):
        title = "Operating KPIs" if i == 0 else "Operating KPIs (cont.)"
        content_slide(prs, title, grp, size=16)

    # Plan
    divider(prs, "The 30 / 60 / 90 plan")
    if deck.get("accountability"):
        content_slide(prs, "Who owns what", deck["accountability"], size=16)
    titles = phase_titles(cfg.get("plan_start", ""))
    for ph, head in zip(deck.get("phases", []), titles):
        full = f"{head.split(' (')[0]}: {ph['title']}" if ph.get("title") else head.split(" (")[0]
        # include date range in title
        date_part = head[head.find("("):] if "(" in head else ""
        content_slide(prs, f"{full} {date_part}".strip(), ph.get("bullets", []), size=16)

    # Ideas
    if deck.get("ideas"):
        content_slide(prs, "Ideas to pursue", deck["ideas"], size=20)

    # Closer
    close = duplicate_slide(prs, SRC_COVER)
    set_ph_lines(close, 0, ["Thank you"])
    set_ph_lines(close, 11, [""])
    set_ph_lines(close, 10, ["acceleratetech.au", "1300 254 321"])

    # delete original example slides
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst)[:n_orig]:
        prs.part.drop_rel(sldId.get(qn('r:id')))
        sldIdLst.remove(sldId)

    prs.save(out_path)
    print("Saved:", out_path, "| slides:", len(prs.slides._sldIdLst))


def main():
    if len(sys.argv) < 2:
        print("Usage: python build_deck.py <content.json>"); sys.exit(1)
    with open(sys.argv[1], encoding="utf-8") as f:
        content = json.load(f)
    cfg = content.get("config", {})
    template = cfg.get("deck_template", DEFAULT_TEMPLATE)
    out_dir = cfg.get("output_dir", ".")
    stem = cfg.get("file_stem", "Operations Deep Dive")
    out = os.path.join(out_dir, f"{cfg.get('deck_name', 'Operations Deep Dive - Readout')}.pptx")
    build(content, out, template)


if __name__ == "__main__":
    main()
